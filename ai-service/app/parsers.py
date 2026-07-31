"""문서 → 평문 텍스트 추출.

지원: pdf, docx, hwpx, txt, md
HWPX는 zip + XML 구조라 외부 라이브러리 없이 직접 파싱한다.
(구형 .hwp는 바이너리 포맷이라 별도 라이브러리가 필요 — W2 이후 과제)
"""
from __future__ import annotations

import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path


class ParseError(Exception):
    """사용자에게 그대로 보여줄 수 있는 파싱 실패 사유."""


SUPPORTED = {"pdf", "docx", "hwpx", "txt", "md"}


def detect_type(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext == "hwp":
        raise ParseError("구버전 .hwp는 아직 지원하지 않습니다. 한글에서 .hwpx로 저장 후 올려주세요.")
    if ext not in SUPPORTED:
        raise ParseError(f"지원하지 않는 형식입니다: .{ext} (지원: {', '.join(sorted(SUPPORTED))})")
    return ext


# ── 개별 파서 ────────────────────────────────────────────────────────

def _parse_pdf(data: bytes) -> str:
    try:
        import pymupdf  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise ParseError("PDF 파서(pymupdf)가 설치되지 않았습니다.") from e

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise ParseError("PDF를 열 수 없습니다. 암호가 걸린 파일인지 확인해주세요.") from e

    if doc.needs_pass:
        raise ParseError("암호가 걸린 PDF는 지원하지 않습니다.")

    parts = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            parts.append(text)
    doc.close()
    return "\n\n".join(parts)


def _parse_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise ParseError("DOCX 파서(python-docx)가 설치되지 않았습니다.") from e

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError("DOCX 파일을 열 수 없습니다.") from e

    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    # 표 안의 텍스트도 놓치지 않는다
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# HWPX 본문 텍스트는 <hp:t> 요소 안에 들어 있다.
_HWPX_TEXT_TAG = "}t"          # 네임스페이스 무시하고 로컬명으로 비교
_HWPX_PARA_TAG = "}p"


def _parse_hwpx(data: bytes) -> str:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ParseError("HWPX 파일이 손상되었거나 형식이 올바르지 않습니다.") from e

    # Contents/section0.xml, section1.xml ... 순서대로
    sections = sorted(
        n for n in zf.namelist()
        if n.startswith("Contents/section") and n.endswith(".xml")
    )
    if not sections:
        raise ParseError("HWPX 본문(Contents/section*.xml)을 찾지 못했습니다.")

    paragraphs: list[str] = []
    for name in sections:
        try:
            root = ET.fromstring(zf.read(name))
        except ET.ParseError:
            continue
        # 문단(<hp:p>) 단위로 <hp:t> 텍스트를 이어 붙인다
        for para in root.iter():
            if not para.tag.endswith(_HWPX_PARA_TAG):
                continue
            buf = [
                node.text for node in para.iter()
                if node.tag.endswith(_HWPX_TEXT_TAG) and node.text
            ]
            line = "".join(buf).strip()
            if line:
                paragraphs.append(line)

    if not paragraphs:
        raise ParseError("HWPX에서 텍스트를 추출하지 못했습니다. 이미지로만 된 문서일 수 있습니다.")
    return "\n".join(paragraphs)


def _parse_text(data: bytes) -> str:
    for enc in ("utf-8", "cp949", "euc-kr"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ParseError("텍스트 인코딩을 인식하지 못했습니다. UTF-8로 저장 후 올려주세요.")


_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "hwpx": _parse_hwpx,
    "txt": _parse_text,
    "md": _parse_text,
}


def normalize(text: str) -> str:
    """공백/개행 정리. 청킹 품질에 직접 영향을 준다."""
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """(file_type, 정규화된 텍스트) 반환."""
    ftype = detect_type(filename)
    text = normalize(_PARSERS[ftype](data))
    if not text:
        raise ParseError("문서에서 텍스트를 찾지 못했습니다.")
    return ftype, text
